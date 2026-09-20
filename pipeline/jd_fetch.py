"""
Stage 3 — JD fetch: split batches, fetch URLs (best-effort), merge into Excel + job.json.

Fetch uses urllib + basic HTML→text distillation. CapSolver/browser not required here.
Failed URLs are marked FETCH_FAILED so the resume stage can still proceed with search.
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import re
import sys
import urllib.error
import urllib.request
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
KEYS_IN = ("row", "num", "company", "title", "jobright", "original")

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
)


class _TextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript", "svg"):
            self._skip += 1

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript", "svg") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if self._skip:
            return
        t = data.strip()
        if t:
            self.parts.append(t)


def html_to_text(html: str) -> str:
    p = _TextExtractor()
    try:
        p.feed(html)
    except Exception:
        pass
    raw = "\n".join(p.parts)
    raw = re.sub(r"\n{3,}", "\n\n", raw)
    return raw.strip()


def fetch_url(url: str, timeout: int = 25) -> str | None:
    if not url or not url.startswith("http"):
        return None
    if "token=" in url.lower():
        return None  # signed URLs often rejected
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            ctype = resp.headers.get("Content-Type", "")
            data = resp.read()
            charset = "utf-8"
            if "charset=" in ctype:
                charset = ctype.split("charset=")[-1].split(";")[0].strip() or "utf-8"
            html = data.decode(charset, errors="replace")
            return html_to_text(html)
    except Exception:
        return None


def distill_jd(text: str, company: str, title: str, max_chars: int = 3500) -> str:
    if not text:
        return "FETCH_FAILED"
    # Drop common junk
    drop = re.compile(
        r"(h-?1b|sponsorship stats|crunchbase|funding|leadership team|"
        r"cookie policy|privacy policy|equal opportunity employer)",
        re.I,
    )
    lines = [ln for ln in text.splitlines() if not drop.search(ln)]
    body = "\n".join(lines).strip()
    if len(body) < 80:
        return "FETCH_FAILED"
    body = body[:max_chars]
    header = f"Title: {title}\nCompany: {company}\n\n"
    return (header + body)[:max_chars]


def _upd(root: Path, day: str) -> Path:
    return root / "Input" / day / "updated"


def _load_jobs(root: Path, day: str):
    path = _upd(root, day) / "jobs_full.json"
    if not path.is_file():
        sys.exit(f"missing jobs_full.json: {path} (run addresses stage first)")
    return json.loads(path.read_text(encoding="utf-8")), path


def cmd_split(root: Path, day: str, n: int = 8) -> dict[str, Any]:
    jobs, src = _load_jobs(root, day)
    n = max(1, int(n))
    bdir = _upd(root, day) / "jd_batches"
    bdir.mkdir(parents=True, exist_ok=True)
    for old in bdir.glob("in_*.json"):
        old.unlink()
    buckets: list[list] = [[] for _ in range(n)]
    for i, j in enumerate(jobs):
        buckets[i % n].append({k: j.get(k, "") for k in KEYS_IN})
    written = 0
    for i, chunk in enumerate(buckets):
        if not chunk and i >= len(jobs):
            continue
        out = bdir / f"in_{i}.json"
        out.write_text(json.dumps(chunk, indent=1, ensure_ascii=False), encoding="utf-8")
        written += 1
    return {"jobs": len(jobs), "batches": written, "dir": str(bdir), "source": str(src)}


def cmd_fetch(root: Path, day: str, batch_idx: int | None = None, force: bool = False) -> dict[str, Any]:
    """Fetch JDs for in_*.json batches and write out_*.json.

    Idempotent: skips batches whose out_*.json already exists unless force=True.
    """
    bdir = _upd(root, day) / "jd_batches"
    ins = sorted(bdir.glob("in_*.json"))
    if batch_idx is not None:
        ins = [p for p in ins if p.name == f"in_{batch_idx}.json"]
    if not ins:
        sys.exit(f"no in_*.json in {bdir} — run split first")
    total = ok = fail = skipped = 0
    for inp in ins:
        outp = bdir / inp.name.replace("in_", "out_")
        if outp.is_file() and not force:
            skipped += 1
            continue
        data = json.loads(inp.read_text(encoding="utf-8"))
        out_list = []
        for ent in data:
            total += 1
            jd = None
            for url_key in ("jobright", "original"):
                txt = fetch_url(str(ent.get(url_key) or ""))
                if txt:
                    jd = distill_jd(txt, ent.get("company", ""), ent.get("title", ""))
                    if jd != "FETCH_FAILED":
                        break
            if not jd or jd == "FETCH_FAILED":
                jd = "FETCH_FAILED"
                fail += 1
            else:
                ok += 1
            out_list.append({"row": ent.get("row"), "num": ent.get("num"), "jd": jd})
        outp.write_text(json.dumps(out_list, indent=1, ensure_ascii=False), encoding="utf-8")
    return {"fetched_ok": ok, "fetch_failed": fail, "total": total, "batches_skipped": skipped}


def cmd_merge(root: Path, day: str) -> dict[str, Any]:
    import openpyxl
    from openpyxl.styles import Alignment

    jobs, _ = _load_jobs(root, day)
    upd = _upd(root, day)
    bdir = upd / "jd_batches"
    out_files = sorted(bdir.glob("out_*.json"))
    if not out_files:
        sys.exit(f"no out_*.json in {bdir}")

    by_key = {}
    for fp in out_files:
        data = json.loads(fp.read_text(encoding="utf-8"))
        for ent in data:
            if isinstance(ent, dict):
                by_key[(ent.get("row"), ent.get("num"))] = ent.get("jd", "")

    by_num = {}
    for (r, n), v in by_key.items():
        by_num.setdefault(n, v)

    cand = [p for p in upd.glob("*_updated.xlsx") if not p.name.startswith("~$")]
    if not cand:
        sys.exit(f"no *_updated.xlsx in {upd}")
    xp = cand[0]
    wb = openpyxl.load_workbook(xp)
    ws = wb[wb.sheetnames[0]]
    cols = (jobs[0].get("cols") or {}) if jobs else {}
    if cols.get("jd"):
        c_jd = cols["jd"]
    else:
        hdr = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
        if "jd" not in hdr:
            sys.exit("Excel has no 'jd' column")
        c_jd = hdr.index("jd") + 1
    wrap = Alignment(wrap_text=True, vertical="top")
    out_root = root / "Resume" / "output" / f"date_{day}"
    merged, missing = 0, []
    for j in jobs:
        key = (j["row"], j["num"])
        if key in by_key:
            jd = by_key[key]
        elif j["num"] in by_num:
            jd = by_num[j["num"]]
        else:
            missing.append(j["num"])
            continue
        ws.cell(row=j["row"], column=c_jd, value=jd).alignment = wrap
        folder = j.get("folder")
        if folder:
            jp = out_root / folder / "job.json"
            if jp.is_file():
                meta = json.loads(jp.read_text(encoding="utf-8"))
                meta["jd"] = jd
                jp.write_text(json.dumps(meta, indent=1, ensure_ascii=False), encoding="utf-8")
        merged += 1
    try:
        wb.save(xp)
    except PermissionError:
        sys.exit(f"PermissionError: close '{xp.name}' and rerun")
    return {
        "merged": merged,
        "total": len(jobs),
        "missing": missing,
        "excel": str(xp),
    }


def main(argv=None):
    ap = argparse.ArgumentParser(description="JD fetch split/fetch/merge")
    ap.add_argument("cmd", choices=("split", "fetch", "merge"))
    ap.add_argument("day", help="M_D_YYYY")
    ap.add_argument("--n", type=int, default=8)
    ap.add_argument("--batch", type=int, default=None)
    ap.add_argument("--root", default=str(ROOT / "data"))
    args = ap.parse_args(argv)
    root = Path(args.root)
    if args.cmd == "split":
        print(json.dumps(cmd_split(root, args.day, args.n), indent=2))
    elif args.cmd == "fetch":
        print(json.dumps(cmd_fetch(root, args.day, args.batch, force=getattr(args, "force", False)), indent=2))
    else:
        r = cmd_merge(root, args.day)
        print(json.dumps(r, indent=2))
        if r["missing"]:
            sys.exit(1)


if __name__ == "__main__":
    main()
