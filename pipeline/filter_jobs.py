"""
Stage 2 - filter / sort / split jobs Excel by role family.

Ported from job-hunt/filter/process_jobs.py (generic paths + config).
"""
from __future__ import annotations

import json
import os
import re
import shutil
import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
CONFIG_DIR = ROOT / "config"
DEFAULT_RULES = CONFIG_DIR / "role_rules.example.json"
DEFAULT_EXCLUDE = CONFIG_DIR / "exclude_companies.example.txt"

MAX_YEARS = 5
UNKNOWN_EXP = 0

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF")


def load_rules(path: str | Path | None = None) -> dict:
    p = Path(path) if path else DEFAULT_RULES
    data_copy = ROOT / "data" / "role_rules.json"
    if data_copy.is_file():
        p = data_copy
    with open(p, encoding="utf-8") as f:
        return json.load(f)


def load_excludes(path: str | Path | None = None) -> list[str]:
    p = Path(path) if path else DEFAULT_EXCLUDE
    data_copy = ROOT / "data" / "exclude_companies.txt"
    if data_copy.is_file():
        p = data_copy
    out: list[str] = []
    if not p.is_file():
        return out
    for line in p.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#"):
            out.append(line)
    return out


def parse_years(value) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    s = str(value).strip().lower()
    if not s or s in ("n/a", "na", "none", "-", "not specified"):
        return None
    if "entry" in s or "intern" in s or "new grad" in s or "graduate" in s:
        return 0
    nums = [int(n) for n in re.findall(r"\d+", s)]
    if not nums:
        return None
    return max(nums)


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def classify(title, rules) -> str:
    t = " " + norm(title) + " "
    for rule in rules["rules"]:
        for token in rule["any"]:
            if token in t:
                return rule["label"]
    return rules.get("fallback_label", "Other Roles")


def safe_name(s) -> str:
    return re.sub(r"[^A-Za-z0-9 _.-]", "", str(s)).strip().replace(" ", "_")[:60] or "Unknown"


def read_rows(path: str):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        import csv
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            data = list(csv.reader(f))
        if not data:
            return [], []
        return data[0], data[1:]

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    headers, rows = None, []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        try:
            first = next(it)
        except StopIteration:
            continue
        head = [str(h).strip() if h is not None else "" for h in first]
        low = [norm(h) for h in head]
        has_title = any("job title" in h or h in ("title", "position", "role") for h in low)
        has_comp = any("company" in h or h == "employer" for h in low)
        if not (has_title and has_comp):
            continue
        if headers is None:
            headers = head
        for r in it:
            if any(c not in (None, "") for c in r):
                rows.append(list(r))
    return headers or [], rows


def col_index(headers, *names):
    low = [norm(h) for h in headers]
    for want in names:
        w = norm(want)
        for i, h in enumerate(low):
            if h == w:
                return i
    for want in names:
        w = norm(want)
        for i, h in enumerate(low):
            if w and w in h:
                return i
    return None


def write_sheet(path, headers, rows, extra_headers):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Jobs"
    full = list(headers) + extra_headers
    ws.append(full)
    for c in range(1, len(full) + 1):
        cell = ws.cell(1, c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(list(r))
    widths: dict[int, int] = {}
    for r in ws.iter_rows(values_only=True):
        for i, v in enumerate(r, 1):
            widths[i] = min(max(widths.get(i, 10), len(str(v or "")) + 2), 55)
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def process(
    inp: str,
    outdir: str | None = None,
    tag: str | None = None,
    max_years: int = MAX_YEARS,
    rules_path: str | None = None,
    exclude_path: str | None = None,
) -> dict[str, Any]:
    rules = load_rules(rules_path)
    excludes = load_excludes(exclude_path)
    excl_norm = [norm(e) for e in excludes if norm(e)]

    headers, rows = read_rows(inp)
    if not headers:
        raise SystemExit("Could not read any rows from " + inp)

    i_title = col_index(headers, "Job Title", "Title", "Position", "Role")
    i_comp = col_index(headers, "Company", "Company Name", "Employer")
    i_exp = col_index(headers, "Experience", "Years of Experience", "Exp", "Experience Required")
    i_match = col_index(headers, "Match %", "Match", "Score")
    if i_title is None or i_comp is None:
        raise SystemExit("Missing a Job Title / Company column. Found: %s" % headers)

    if tag is None:
        t = datetime.date.today()
        tag = f"{t.month}_{t.day}_{t.year}"
    outdir = outdir or str(ROOT / "data" / "Output" / tag)
    if os.path.isdir(outdir):
        shutil.rmtree(outdir)
    os.makedirs(outdir, exist_ok=True)

    groups: dict[str, list] = {}
    dropped_exp = dropped_comp = dupes = 0
    removed_rows: list = []
    seen: set = set()

    for row in rows:
        row = list(row) + [None] * (len(headers) - len(row))
        title = row[i_title]
        comp = row[i_comp]
        if not title and not comp:
            continue

        key = (norm(title), norm(comp))
        if key in seen:
            dupes += 1
            continue
        seen.add(key)

        hit = next((e for e in excl_norm if e and e in norm(comp)), None)
        if hit:
            dropped_comp += 1
            removed_rows.append(list(row) + ["excluded company: %s" % comp])
            continue

        yrs_raw = parse_years(row[i_exp]) if i_exp is not None else None
        yrs = UNKNOWN_EXP if yrs_raw is None else yrs_raw
        if yrs > max_years:
            dropped_exp += 1
            removed_rows.append(list(row) + ["%s+ years experience" % yrs])
            continue

        label = classify(title, rules)
        try:
            match = (
                float(re.sub(r"[^\d.]", "", str(row[i_match])))
                if i_match is not None and row[i_match] not in (None, "")
                else -1
            )
        except ValueError:
            match = -1

        groups.setdefault(label, []).append(
            (yrs, -match, list(row) + [yrs, "yes" if yrs_raw is None else "no"])
        )

    extra = ["Experience (years)", "Experience Missing"]
    files, summary = [], []
    for label in sorted(groups, key=lambda k: (-len(groups[k]), k)):
        items = sorted(groups[label], key=lambda x: (x[0], x[1]))
        out_rows = [it[2] for it in items]
        fname = "%s__%d_jobs.xlsx" % (safe_name(label), len(out_rows))
        fpath = os.path.join(outdir, fname)
        write_sheet(fpath, headers, out_rows, extra)
        files.append(fpath)
        summary.append({"role": label, "count": len(out_rows), "file": fname})

    if removed_rows:
        rpath = os.path.join(outdir, "_REMOVED__%d_jobs.xlsx" % len(removed_rows))
        write_sheet(rpath, headers, removed_rows, ["Removed Because"])
        files.append(rpath)

    kept = sum(s["count"] for s in summary)
    result = {
        "input": inp,
        "output_dir": outdir,
        "tag": tag,
        "total_rows": len(rows),
        "duplicates_skipped": dupes,
        "dropped_experience_6plus": dropped_exp,
        "dropped_excluded_company": dropped_comp,
        "kept": kept,
        "groups": summary,
        "files": files,
    }
    return result


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        raise SystemExit("usage: filter_jobs.py <input.xlsx> [output_dir] [tag]")
    r = process(
        sys.argv[1],
        sys.argv[2] if len(sys.argv) > 2 else None,
        sys.argv[3] if len(sys.argv) > 3 else None,
    )
    print(json.dumps(r, indent=2))
